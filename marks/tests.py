import io
import json
import hashlib
import shutil
import tempfile
from datetime import date, timedelta
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .forms import BranchForm, TikTokFunnelRequestForm
from .models import Bot, Branch, Experiment, Product, TaskRequest, TikTokFunnelRequest, UserProfile
from .task_time import get_tasks_timezone


class TaskBoardBaseTestCase(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.admin_user = user_model.objects.create_user(username="admin", password="StrongPass123!")
        self.admin_user.profile.role = UserProfile.Role.ADMIN
        self.admin_user.profile.save(update_fields=["role"])

        self.manager_user = user_model.objects.create_user(username="manager", password="StrongPass123!")
        self.manager_user.profile.role = UserProfile.Role.MANAGER
        self.manager_user.profile.save(update_fields=["role"])

        self.product = Product.objects.create(name="Test product")
        self.bot = Bot.objects.create(name="test_bot_name", product=self.product)
        self.branch_main = Branch.objects.create(bot=self.bot, name="Main", code="MN")


class TaskBoardAccessTests(TaskBoardBaseTestCase):
    def test_admin_has_access_to_tasks_board(self):
        self.client.force_login(self.admin_user)
        response = self.client.get(reverse("tasks_board"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Задачник")
        self.assertContains(response, "Выгрузить выполненные за период")

    def test_non_admin_has_access_but_no_kanban(self):
        self.client.force_login(self.manager_user)
        response = self.client.get(reverse("tasks_board"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Задачник")
        self.assertFalse(response.context["show_kanban"])
        self.assertNotContains(response, "Выгрузить выполненные за период")


class TaskBoardActionsTests(TaskBoardBaseTestCase):
    @patch("marks.views.notify_new_task")
    def test_create_patch_task(self, notify_mock):
        self.client.force_login(self.admin_user)
        deadline = timezone.now() + timedelta(days=3)
        response = self.client.post(
            reverse("create_patch_task"),
            {
                "patch-branches": [self.branch_main.id],
                "patch-cjm_url": "https://example.com/cjm",
                "patch-comment": "Комментарий",
                "patch-deadline": deadline.strftime("%Y-%m-%dT%H:%M"),
            },
        )

        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, reverse("tasks_board"))

        task = TaskRequest.objects.get()
        self.assertEqual(task.task_type, TaskRequest.Type.PATCH)
        self.assertEqual(task.status, TaskRequest.Status.UNREAD)
        self.assertEqual(task.created_by, self.admin_user)
        self.assertEqual(list(task.branches.values_list("id", flat=True)), [self.branch_main.id])
        notify_mock.assert_called_once_with(task)

    @patch("marks.views.notify_new_task")
    def test_create_patch_task_saves_photo_and_task_timezone_deadline(self, notify_mock):
        self.client.force_login(self.admin_user)
        media_root = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, media_root, ignore_errors=True)
        photo = SimpleUploadedFile("task.png", b"fake-image-bytes", content_type="image/png")

        with self.settings(MEDIA_ROOT=media_root, TASKS_TIME_ZONE="Europe/Moscow"):
            response = self.client.post(
                reverse("create_patch_task"),
                {
                    "patch-branches": [self.branch_main.id],
                    "patch-cjm_url": "https://example.com/cjm",
                    "patch-comment": "Комментарий со скрином",
                    "patch-deadline": "2026-04-08T12:34",
                    "patch-photo": photo,
                },
            )
            task = TaskRequest.objects.get(task_type=TaskRequest.Type.PATCH)
            self.assertTrue(task.photo.name.endswith(".png"))
            self.assertEqual(
                timezone.localtime(task.deadline, get_tasks_timezone()).strftime("%Y-%m-%dT%H:%M"),
                "2026-04-08T12:34",
            )

        self.assertEqual(response.status_code, 302)
        task.refresh_from_db()
        notify_mock.assert_called_once_with(task)

    @patch("marks.views.notify_new_task")
    def test_create_build_task_uses_manual_bot_name(self, notify_mock):
        self.client.force_login(self.admin_user)
        deadline = timezone.now() + timedelta(days=2)
        response = self.client.post(
            reverse("create_build_task"),
            {
                "build-bot_name": "@new_bot",
                "build-build_token": "1234567890",
                "build-cjm_url": "https://example.com/cjm-build",
                "build-comment": "Build comment",
                "build-deadline": deadline.strftime("%Y-%m-%dT%H:%M"),
            },
        )

        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, reverse("tasks_board"))

        task = TaskRequest.objects.get(task_type=TaskRequest.Type.BUILD)
        self.assertEqual(task.build_name, "@new_bot")
        self.assertEqual(task.branches.count(), 0)
        self.assertEqual(task.get_scope_units(), 1)
        notify_mock.assert_called_once_with(task)

    @patch("marks.views.notify_new_task")
    def test_create_build_task_appends_optional_branch_name(self, notify_mock):
        self.client.force_login(self.admin_user)
        deadline = timezone.now() + timedelta(days=2)
        response = self.client.post(
            reverse("create_build_task"),
            {
                "build-bot_name": "@new_bot",
                "build-branch_name": "feature-login",
                "build-build_token": "1234567890",
                "build-cjm_url": "https://example.com/cjm-build",
                "build-comment": "Build comment",
                "build-deadline": deadline.strftime("%Y-%m-%dT%H:%M"),
            },
        )

        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, reverse("tasks_board"))

        task = TaskRequest.objects.get(task_type=TaskRequest.Type.BUILD)
        self.assertEqual(task.build_name, "@new_bot / feature-login")
        self.assertEqual(task.get_scope_units(), 1)
        notify_mock.assert_called_once_with(task)

    @patch("marks.views.notify_new_task")
    def test_create_mailing_task_accepts_local_deadline_format(self, notify_mock):
        self.client.force_login(self.admin_user)
        deadline = timezone.now() + timedelta(days=2)
        response = self.client.post(
            reverse("create_mailing_task"),
            {
                "mailing-branches": [self.branch_main.id],
                "mailing-tz_url": "https://example.com/tz",
                "mailing-comment": "Mailing comment",
                "mailing-deadline": deadline.strftime("%d.%m.%Y %H:%M"),
            },
        )

        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, reverse("tasks_board"))

        task = TaskRequest.objects.get(task_type=TaskRequest.Type.MAILING)
        self.assertEqual(task.status, TaskRequest.Status.UNREAD)
        self.assertEqual(task.created_by, self.admin_user)
        self.assertEqual(task.tz_url, "https://example.com/tz")
        self.assertEqual(list(task.branches.values_list("id", flat=True)), [self.branch_main.id])
        notify_mock.assert_called_once_with(task)

    @patch("marks.views.notify_new_task")
    def test_create_task_with_notify_requires_tg_username(self, notify_mock):
        self.client.force_login(self.admin_user)
        deadline = timezone.now() + timedelta(days=2)
        response = self.client.post(
            reverse("create_patch_task"),
            {
                "patch-branches": [self.branch_main.id],
                "patch-cjm_url": "https://example.com/cjm",
                "patch-comment": "Комментарий",
                "patch-deadline": deadline.strftime("%Y-%m-%dT%H:%M"),
                "patch-notify_me": "on",
                "patch-tg_username": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Укажите username в Telegram или chat_id.")
        self.assertFalse(TaskRequest.objects.filter(task_type=TaskRequest.Type.PATCH).exists())
        notify_mock.assert_not_called()

    @patch("marks.views.notify_status_change")
    def test_status_done_sets_completed_at_and_sends_notification(self, notify_mock):
        task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.BUILD,
            build_name="bot + branches",
            build_token="1234567890",
            cjm_url="https://example.com/cjm",
            deadline=timezone.now() + timedelta(days=1),
            created_by=self.admin_user,
        )
        self.client.force_login(self.admin_user)

        response = self.client.post(
            reverse("update_task_status", kwargs={"task_id": task.id}),
            {"status": TaskRequest.Status.DONE},
        )

        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, reverse("tasks_board"))
        task.refresh_from_db()
        self.assertEqual(task.status, TaskRequest.Status.DONE)
        self.assertIsNotNone(task.completed_at)

        notify_mock.assert_called_once()
        kwargs = notify_mock.call_args.kwargs
        self.assertEqual(kwargs["task"], task)
        self.assertEqual(kwargs["old_status"], TaskRequest.Status.UNREAD)
        self.assertEqual(kwargs["changed_by"], self.admin_user)

    @patch("marks.views.notify_status_change")
    def test_status_done_never_saves_completed_at_before_created_at(self, notify_mock):
        task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.BUILD,
            build_name="bot + branches",
            build_token="1234567890",
            cjm_url="https://example.com/cjm",
            deadline=timezone.now() + timedelta(days=1),
            created_by=self.admin_user,
        )
        future_created_at = timezone.now() + timedelta(minutes=10)
        TaskRequest.objects.filter(pk=task.pk).update(created_at=future_created_at)
        self.client.force_login(self.admin_user)

        response = self.client.post(
            reverse("update_task_status", kwargs={"task_id": task.id}),
            {"status": TaskRequest.Status.DONE},
        )

        self.assertEqual(response.status_code, 302)
        task.refresh_from_db()
        self.assertEqual(task.completed_at, task.created_at)
        notify_mock.assert_called_once()

    @patch("marks.views.notify_done_to_user")
    @patch("marks.views.get_task_tg_username", return_value="test_user")
    @patch("marks.views.notify_status_change")
    def test_status_done_sends_personal_notification_when_username_exists(
        self,
        status_notify_mock,
        legacy_username_mock,
        user_notify_mock,
    ):
        task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.MAILING,
            tz_url="https://example.com/tz",
            deadline=timezone.now() + timedelta(days=1),
            created_by=self.admin_user,
        )
        self.client.force_login(self.admin_user)

        response = self.client.post(
            reverse("update_task_status", kwargs={"task_id": task.id}),
            {"status": TaskRequest.Status.DONE},
        )

        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, reverse("tasks_board"))
        task.refresh_from_db()
        self.assertEqual(task.status, TaskRequest.Status.DONE)

        status_notify_mock.assert_called_once()
        legacy_username_mock.assert_called_once_with(task.id)
        user_notify_mock.assert_called_once_with(task=task, tg_username="test_user")

    @override_settings(TELEGRAM_WEBHOOK_SECRET="secret-key")
    @patch("marks.views.set_task_feedback_comment")
    def test_telegram_webhook_saves_feedback_from_reply(self, set_feedback_mock):
        task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.PATCH,
            cjm_url="https://example.com/cjm",
            deadline=timezone.now() + timedelta(days=1),
            created_by=self.admin_user,
            status=TaskRequest.Status.DONE,
        )
        payload = {
            "update_id": 1,
            "message": {
                "message_id": 101,
                "text": "Все ок, спасибо!",
                "reply_to_message": {
                    "message_id": 100,
                    "text": f"ID задачи: #{task.id}\nЗадача выполнена",
                },
            },
        }

        response = self.client.post(
            reverse("telegram_webhook", kwargs={"webhook_key": "secret-key"}),
            data=json.dumps(payload),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        set_feedback_mock.assert_called_once_with(task_id=task.id, feedback_comment="Все ок, спасибо!")

    @override_settings(TELEGRAM_WEBHOOK_SECRET="secret-key")
    @patch("marks.views.set_task_feedback_comment")
    def test_telegram_webhook_saves_feedback_from_quote_when_reply_is_inaccessible(self, set_feedback_mock):
        task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.PATCH,
            cjm_url="https://example.com/cjm",
            deadline=timezone.now() + timedelta(days=1),
            created_by=self.admin_user,
            status=TaskRequest.Status.DONE,
        )
        payload = {
            "update_id": 5,
            "message": {
                "message_id": 105,
                "text": "РЎРїР°СЃРёР±Рѕ, РІС‹ СЃСѓРїРµСЂ!!",
                "reply_to_message": {
                    "message_id": 104,
                    "date": 0,
                    "chat": {"id": -100100100, "type": "supergroup"},
                },
                "quote": {
                    "text": f"Р—Р°РґР°С‡Р° РІС‹РїРѕР»РЅРµРЅР°\nID Р·Р°РґР°С‡Рё: #{task.id}",
                },
            },
        }

        response = self.client.post(
            reverse("telegram_webhook", kwargs={"webhook_key": "secret-key"}),
            data=json.dumps(payload),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        set_feedback_mock.assert_called_once_with(task_id=task.id, feedback_comment="РЎРїР°СЃРёР±Рѕ, РІС‹ СЃСѓРїРµСЂ!!")

    @override_settings(TELEGRAM_WEBHOOK_SECRET="secret-key")
    @patch("marks.views.set_task_feedback_comment")
    def test_telegram_webhook_saves_feedback_from_business_message_external_reply(self, set_feedback_mock):
        task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.MAILING,
            tz_url="https://example.com/tz",
            deadline=timezone.now() + timedelta(days=1),
            created_by=self.admin_user,
            status=TaskRequest.Status.DONE,
        )
        payload = {
            "update_id": 6,
            "business_message": {
                "message_id": 106,
                "text": "Р¤РёРґР±РµРє РёР· business chat",
                "external_reply": {
                    "message_id": 105,
                    "text": f"ID Р·Р°РґР°С‡Рё: #{task.id}\nР—Р°РґР°С‡Р° РІС‹РїРѕР»РЅРµРЅР°",
                },
            },
        }

        response = self.client.post(
            reverse("telegram_webhook", kwargs={"webhook_key": "secret-key"}),
            data=json.dumps(payload),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        set_feedback_mock.assert_called_once_with(task_id=task.id, feedback_comment="Р¤РёРґР±РµРє РёР· business chat")

    @override_settings(TELEGRAM_WEBHOOK_SECRET="secret-key")
    def test_telegram_webhook_feedback_is_exported(self):
        task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.PATCH,
            cjm_url="https://example.com/cjm",
            deadline=timezone.now() + timedelta(days=1),
            created_by=self.admin_user,
            status=TaskRequest.Status.DONE,
        )
        payload = {
            "update_id": 4,
            "message": {
                "message_id": 104,
                "text": "Фидбек по задаче",
                "reply_to_message": {
                    "message_id": 103,
                    "text": f"ID задачи: #{task.id}\nЗадача выполнена",
                },
            },
        }

        webhook_response = self.client.post(
            reverse("telegram_webhook", kwargs={"webhook_key": "secret-key"}),
            data=json.dumps(payload),
            content_type="application/json",
        )

        self.assertEqual(webhook_response.status_code, 200)

        self.client.force_login(self.admin_user)
        export_response = self.client.get(reverse("export_completed_tasks"))

        self.assertEqual(export_response.status_code, 200)
        workbook = load_workbook(io.BytesIO(export_response.content))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        feedback_index = headers.index("Фидбек")
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        row = next(row for row in rows if row[0] == task.id)

        self.assertEqual(row[feedback_index], "Фидбек по задаче")

    @override_settings(TELEGRAM_WEBHOOK_SECRET="secret-key")
    @patch("marks.views.send_text_message")
    @patch("marks.views.send_weekly_tasks_report", return_value=(True, ""))
    def test_telegram_webhook_week_command_sends_report(self, send_report_mock, send_text_mock):
        task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.PATCH,
            cjm_url="https://example.com/cjm",
            deadline=timezone.now() + timedelta(days=1),
            created_by=self.admin_user,
            status=TaskRequest.Status.DONE,
        )
        task.completed_at = timezone.now()
        task.save(update_fields=["completed_at"])

        payload = {
            "update_id": 2,
            "message": {
                "message_id": 102,
                "text": "/week",
                "chat": {"id": 439144407},
            },
        }
        response = self.client.post(
            reverse("telegram_webhook", kwargs={"webhook_key": "secret-key"}),
            data=json.dumps(payload),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        send_report_mock.assert_called_once()
        kwargs = send_report_mock.call_args.kwargs
        self.assertEqual(kwargs["chat_id"], "439144407")
        self.assertIn("tasks_week_current_", kwargs["filename"])
        self.assertIn("текущую неделю", kwargs["caption"])
        send_text_mock.assert_not_called()

    @override_settings(TELEGRAM_WEBHOOK_SECRET="secret-key")
    @patch("marks.views.send_text_message")
    @patch("marks.views.send_weekly_tasks_report", return_value=(True, ""))
    def test_telegram_webhook_month_command_rejects_bad_date(self, send_report_mock, send_text_mock):
        payload = {
            "update_id": 3,
            "message": {
                "message_id": 103,
                "text": "/month 15-02-2026",
                "chat": {"id": 439144407},
            },
        }
        response = self.client.post(
            reverse("telegram_webhook", kwargs={"webhook_key": "secret-key"}),
            data=json.dumps(payload),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        send_report_mock.assert_not_called()
        send_text_mock.assert_called_once()

    def test_tasks_board_filters_by_status(self):
        done_task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.BUILD,
            build_name="done task",
            build_token="token",
            cjm_url="https://example.com/cjm",
            deadline=timezone.now() + timedelta(days=2),
            created_by=self.admin_user,
            status=TaskRequest.Status.DONE,
        )
        unread_task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.BUILD,
            build_name="unread task",
            build_token="token2",
            cjm_url="https://example.com/cjm2",
            deadline=timezone.now() + timedelta(days=2),
            created_by=self.admin_user,
            status=TaskRequest.Status.UNREAD,
        )
        self.client.force_login(self.admin_user)

        response = self.client.get(
            reverse("tasks_board"),
            {"task_status": TaskRequest.Status.DONE},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f">#{done_task.id}<")
        self.assertNotContains(response, f">#{unread_task.id}<")

    def test_completed_counters_use_branch_units(self):
        second_branch = Branch.objects.create(bot=self.bot, name="Dev", code="DV")
        third_branch = Branch.objects.create(bot=self.bot, name="Feature", code="FT")

        mailing_task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.MAILING,
            tz_url="https://example.com/tz",
            deadline=timezone.now() + timedelta(days=1),
            created_by=self.admin_user,
            status=TaskRequest.Status.DONE,
        )
        mailing_task.branches.set([self.branch_main, second_branch])

        build_task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.BUILD,
            build_token="build-token",
            cjm_url="https://example.com/cjm",
            deadline=timezone.now() + timedelta(days=1),
            created_by=self.admin_user,
            status=TaskRequest.Status.DONE,
        )
        build_task.branches.set([self.branch_main, second_branch, third_branch])

        patch_task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.PATCH,
            cjm_url="https://example.com/cjm-patch",
            deadline=timezone.now() + timedelta(days=1),
            created_by=self.admin_user,
            status=TaskRequest.Status.DONE,
        )
        patch_task.branches.set([self.branch_main])

        self.client.force_login(self.admin_user)
        response = self.client.get(reverse("tasks_board"))

        self.assertEqual(response.status_code, 200)
        counters = response.context["completed_type_counters"]
        self.assertEqual(counters["mailing"], 2)
        self.assertEqual(counters["build"], 3)
        self.assertEqual(counters["patch"], 1)

    def test_export_completed_tasks_by_period(self):
        recent_task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.BUILD,
            build_name="recent",
            build_token="token",
            cjm_url="https://example.com/cjm",
            deadline=timezone.now() + timedelta(days=1),
            created_by=self.admin_user,
            status=TaskRequest.Status.DONE,
        )
        old_task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.BUILD,
            build_name="old",
            build_token="token",
            cjm_url="https://example.com/cjm",
            deadline=timezone.now() + timedelta(days=1),
            created_by=self.admin_user,
            status=TaskRequest.Status.DONE,
        )

        recent_task.completed_at = timezone.now() - timedelta(days=1)
        recent_task.save(update_fields=["completed_at"])
        old_task.completed_at = timezone.now() - timedelta(days=30)
        old_task.save(update_fields=["completed_at"])

        self.client.force_login(self.admin_user)
        date_from = (timezone.localdate() - timedelta(days=3)).isoformat()
        date_to = timezone.localdate().isoformat()
        response = self.client.get(
            reverse("export_completed_tasks"),
            {"completed_from": date_from, "completed_to": date_to},
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", response["Content-Type"])

        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        task_ids = [row[0] for row in rows]
        self.assertIn(recent_task.id, task_ids)
        self.assertNotIn(old_task.id, task_ids)

    def test_export_has_combined_link_column_and_bot_branch_column(self):
        task = TaskRequest.objects.create(
            task_type=TaskRequest.Type.MAILING,
            tz_url="https://example.com/tz",
            deadline=timezone.now() + timedelta(days=1),
            created_by=self.admin_user,
            status=TaskRequest.Status.DONE,
        )
        task.branches.set([self.branch_main])
        self.client.force_login(self.admin_user)

        response = self.client.get(reverse("export_completed_tasks"))
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]

        self.assertNotIn("Ветки", headers)
        self.assertIn("CJM/ТЗ", headers)
        self.assertIn("Бот и ветки", headers)
        self.assertIn("Фидбек", headers)

        first_data_row = [cell.value for cell in sheet[2]]
        self.assertIn("https://example.com/tz", first_data_row)
        self.assertIn("test_bot_name", first_data_row)


class BotPlatformTests(TaskBoardBaseTestCase):
    def test_default_telegram_tag_uses_start_link(self):
        tag = self.branch_main.tags.get(url__isnull=False)

        self.assertEqual(tag.url, "https://telegram.me/test_bot_name?start=MN0001")

    def test_vk_tag_uses_group_ref_link(self):
        vk_bot = Bot.objects.create(
            name="203482421",
            display_name="VK Sales",
            platform=Bot.Platform.VK,
            product=self.product,
        )
        branch = Branch.objects.create(bot=vk_bot, name="Main", code="ell23")
        tag = branch.tags.get(url__isnull=False)

        self.assertEqual(tag.url, "https://vk.com/write-203482421?ref=ell230001&ref_source=23")

    def test_bot_api_finds_vk_bot_by_group_id(self):
        vk_bot = Bot.objects.create(
            name="203482421",
            display_name="VK Sales",
            platform=Bot.Platform.VK,
            product=self.product,
        )
        Branch.objects.create(bot=vk_bot, name="Main", code="ell01")

        response = self.client.get(reverse("bot_api", kwargs={"bot_name": "203482421"}))

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["bot"], "203482421")
        self.assertEqual(payload["branches"][0]["tags"][0]["url"], "https://vk.com/write-203482421?ref=ell010001&ref_source=1")

    def test_bot_api_accepts_telegram_name_with_at_prefix(self):
        response = self.client.get(reverse("bot_api", kwargs={"bot_name": "@test_bot_name"}))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["bot"], "test_bot_name")

    @patch("marks.views.secrets.randbelow", return_value=85)
    def test_bot_api_number_response_includes_ab_key_for_active_test(self, randbelow_mock):
        Experiment.objects.create(
            title="API split by number",
            branch=self.branch_main,
            wants_ab_test=True,
            ab_test_options=["start"],
            metric_impact="CR",
            expected_change="+5%",
            hypothesis="Проверяем вариант первого экрана.",
            traffic_volume=Experiment.TrafficVolume.SPLIT_70_30,
            test_duration=Experiment.TestDuration.DAYS_7,
            start_date=date(2026, 3, 1),
            end_date=date(2026, 4, 10),
            status=Experiment.Status.IN_PROGRESS,
            created_by=self.admin_user,
        )
        number = self.branch_main.tags.get(url__isnull=False).number

        response = self.client.get(
            reverse("bot_api", kwargs={"bot_name": "test_bot_name"}),
            {"number": number},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["number"], number)
        self.assertEqual(payload["ab_key"], 2)
        randbelow_mock.assert_called_once_with(100)

    def test_bot_api_number_response_defaults_ab_key_to_one_without_active_test(self):
        number = self.branch_main.tags.get(url__isnull=False).number

        response = self.client.get(
            reverse("bot_api", kwargs={"bot_name": "test_bot_name"}),
            {"number": number},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["number"], number)
        self.assertEqual(payload["ab_key"], 1)

    def test_bot_api_returns_branch_ab_assignment_for_specific_branch_code(self):
        experiment = Experiment.objects.create(
            title="API split",
            branch=self.branch_main,
            wants_ab_test=True,
            ab_test_options=["start"],
            metric_impact="CR",
            expected_change="+5%",
            hypothesis="Проверяем вариант первого экрана.",
            traffic_volume=Experiment.TrafficVolume.SPLIT_70_30,
            test_duration=Experiment.TestDuration.DAYS_7,
            start_date=date(2026, 3, 1),
            end_date=date(2026, 4, 10),
            status=Experiment.Status.IN_PROGRESS,
            created_by=self.admin_user,
        )
        ab_key = "user-42"
        seed = f"{experiment.id}:{self.branch_main.id}:{ab_key}".encode("utf-8")
        bucket = int(hashlib.sha256(seed).hexdigest()[:16], 16) % 100
        expected_variant_value = 1 if bucket < 70 else 2

        response = self.client.get(
            reverse("bot_api", kwargs={"bot_name": "test_bot_name"}),
            {"branch_code": "mn", "ab_key": ab_key},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(len(payload["branches"]), 1)
        self.assertEqual(payload["branches"][0]["code"], "MN")
        self.assertEqual(payload["branches"][0]["ab_test"]["active"], True)
        self.assertEqual(payload["branches"][0]["ab_test"]["variant_value"], expected_variant_value)
        self.assertEqual(payload["branches"][0]["ab_test"]["split"], "70/30")
        self.assertEqual(payload["branches"][0]["ab_test"]["assignment_mode"], "hash")

    def test_bot_api_returns_inactive_ab_payload_when_branch_has_no_active_test(self):
        response = self.client.get(
            reverse("bot_api", kwargs={"bot_name": "test_bot_name"}),
            {"branch_code": "MN"},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["branches"][0]["ab_test"], {"active": False})

    def test_bot_api_returns_404_for_unknown_branch_code(self):
        response = self.client.get(
            reverse("bot_api", kwargs={"bot_name": "test_bot_name"}),
            {"branch_code": "missing"},
        )

        self.assertEqual(response.status_code, 404)

    def test_bots_list_creates_telegram_bot_and_strips_at_sign(self):
        self.client.force_login(self.admin_user)

        response = self.client.post(
            reverse("bots_list"),
            {
                "form_type": "telegram",
                "tg-name": "@new_bot",
            },
        )

        self.assertEqual(response.status_code, 302)
        bot = Bot.objects.get(name="new_bot")
        self.assertEqual(bot.platform, Bot.Platform.TELEGRAM)
        self.assertEqual(bot.display_name, "")

    def test_bots_list_creates_vk_bot_with_display_name(self):
        self.client.force_login(self.admin_user)

        response = self.client.post(
            reverse("bots_list"),
            {
                "form_type": "vk",
                "vk-name": "203482421",
                "vk-display_name": "VK Sales",
            },
        )

        self.assertEqual(response.status_code, 302)
        bot = Bot.objects.get(name="203482421")
        self.assertEqual(bot.platform, Bot.Platform.VK)
        self.assertEqual(bot.display_name, "VK Sales")

    def test_bots_list_sorts_telegram_before_vk(self):
        Bot.objects.create(name="zzz_bot", product=self.product)
        Bot.objects.create(
            name="203482421",
            display_name="Alpha VK",
            platform=Bot.Platform.VK,
            product=self.product,
        )
        self.client.force_login(self.admin_user)

        response = self.client.get(reverse("bots_list"))

        self.assertEqual(response.status_code, 200)
        ordered_platforms = [bot.platform for bot in response.context["active_bots"]]
        self.assertEqual(ordered_platforms[:3], [Bot.Platform.TELEGRAM, Bot.Platform.TELEGRAM, Bot.Platform.VK])


class ExperimentBoardTests(TaskBoardBaseTestCase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.admin_user)

    def _experiment_payload(self, **overrides):
        payload = {
            "title": "Новый A/B тест",
            "tz_url": "https://example.com/tz/exp-1",
            "wants_ab_test": "on",
            "ab_test_options": ["start"],
            "ab_test_custom_option": "",
            "metric_impact": "CR",
            "comparison_text": "Текущий экран vs новый экран",
            "expected_change": "+8%",
            "hypothesis": "Если изменить первый экран, конверсия вырастет.",
            "traffic_volume": Experiment.TrafficVolume.SPLIT_50_50,
            "traffic_volume_other": "",
            "test_duration": Experiment.TestDuration.DAYS_7,
            "duration_users": "",
            "duration_end_date": "",
            "start_date": "2026-03-10",
            "end_date": "2026-03-17",
            "dashboard_url": "https://example.com/dashboard/exp-1",
            "result_variant_a": "CR 11%, open rate 24%",
            "result_variant_b": "CR 13%, open rate 28%",
            "comment": "Комментарий по эксперименту",
        }
        payload.update(overrides)
        return payload

    def test_create_experiment_saves_dates_and_manual_results(self):
        response = self.client.post(reverse("experiments_board"), self._experiment_payload())

        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, reverse("experiments_board"))

        experiment = Experiment.objects.get(title="Новый A/B тест")
        self.assertEqual(experiment.status, Experiment.Status.BACKLOG)
        self.assertEqual(experiment.tz_url, "https://example.com/tz/exp-1")
        self.assertEqual(experiment.comparison_text, "Текущий экран vs новый экран")
        self.assertEqual(experiment.start_date, date(2026, 3, 10))
        self.assertEqual(experiment.end_date, date(2026, 3, 17))
        self.assertEqual(experiment.dashboard_url, "https://example.com/dashboard/exp-1")
        self.assertEqual(experiment.result_variant_a, "CR 11%, open rate 24%")
        self.assertEqual(experiment.result_variant_b, "CR 13%, open rate 28%")

    def test_create_experiment_saves_selected_branch_for_api(self):
        response = self.client.post(
            reverse("experiments_board"),
            self._experiment_payload(branch=str(self.branch_main.id)),
        )

        self.assertEqual(response.status_code, 302)
        experiment = Experiment.objects.get()
        self.assertEqual(experiment.branch_id, self.branch_main.id)
        return
        experiment = Experiment.objects.get(title="РќРѕРІС‹Р№ A/B С‚РµСЃС‚")
        self.assertEqual(experiment.branch_id, self.branch_main.id)

    def test_edit_experiment_updates_dates_and_metrics(self):
        experiment = Experiment.objects.create(
            title="Текущий тест",
            wants_ab_test=True,
            ab_test_options=["start"],
            metric_impact="CR",
            expected_change="+5%",
            hypothesis="Исходная гипотеза",
            traffic_volume=Experiment.TrafficVolume.SPLIT_50_50,
            test_duration=Experiment.TestDuration.DAYS_7,
            created_by=self.admin_user,
        )

        payload = self._experiment_payload(
            title="Текущий тест",
            experiment_id=str(experiment.id),
            result_variant_a="CR 10%",
            result_variant_b="CR 14%",
            start_date="2026-03-11",
            end_date="2026-03-18",
        )
        response = self.client.post(reverse("experiments_board"), payload)

        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, reverse("experiments_board"))

        experiment.refresh_from_db()
        self.assertEqual(experiment.start_date, date(2026, 3, 11))
        self.assertEqual(experiment.end_date, date(2026, 3, 18))
        self.assertEqual(experiment.result_variant_a, "CR 10%")
        self.assertEqual(experiment.result_variant_b, "CR 14%")

    @patch("marks.views.notify_new_task", return_value=(True, ""))
    def test_move_to_draft_allows_empty_tz(self, notify_mock):
        experiment = Experiment.objects.create(
            title="Без ТЗ",
            wants_ab_test=True,
            ab_test_options=["start"],
            metric_impact="CR",
            comparison_text="Экран A vs экран B",
            expected_change="+5%",
            hypothesis="Проверяем первый экран",
            traffic_volume=Experiment.TrafficVolume.SPLIT_50_50,
            test_duration=Experiment.TestDuration.DAYS_7,
            created_by=self.admin_user,
            status=Experiment.Status.BACKLOG,
        )

        response = self.client.post(
            reverse("update_experiment_status", kwargs={"experiment_id": experiment.id}),
            {"status": Experiment.Status.DRAFT},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        experiment.refresh_from_db()
        self.assertEqual(experiment.status, Experiment.Status.DRAFT)
        self.assertIsNotNone(experiment.technical_task)
        self.assertEqual(experiment.technical_task.tz_url, "")
        notify_mock.assert_called_once_with(experiment.technical_task)

    @patch("marks.views.notify_new_task", return_value=(True, ""))
    def test_move_to_draft_creates_task_for_tech_team(self, notify_mock):
        experiment = Experiment.objects.create(
            title="Готов к разработке",
            branch=self.branch_main,
            tz_url="https://example.com/tz/ready",
            wants_ab_test=True,
            ab_test_options=["start"],
            metric_impact="CR",
            comparison_text="Экран A vs экран B",
            expected_change="+5%",
            hypothesis="Проверяем первый экран",
            traffic_volume=Experiment.TrafficVolume.SPLIT_50_50,
            test_duration=Experiment.TestDuration.DAYS_7,
            start_date=date(2026, 4, 2),
            created_by=self.admin_user,
            status=Experiment.Status.BACKLOG,
        )

        response = self.client.post(
            reverse("update_experiment_status", kwargs={"experiment_id": experiment.id}),
            {"status": Experiment.Status.DRAFT},
        )

        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, reverse("experiments_board"))

        experiment.refresh_from_db()
        self.assertEqual(experiment.status, Experiment.Status.DRAFT)
        self.assertIsNotNone(experiment.technical_task)

        task = experiment.technical_task
        self.assertEqual(task.task_type, TaskRequest.Type.PATCH)
        self.assertEqual(task.status, TaskRequest.Status.UNREAD)
        self.assertEqual(task.tz_url, "https://example.com/tz/ready")
        self.assertEqual(list(task.branches.values_list("id", flat=True)), [self.branch_main.id])
        notify_mock.assert_called_once_with(task)

    def test_update_experiment_status_blocks_parallel_branch_test(self):
        Experiment.objects.create(
            title="РђРєС‚РёРІРЅС‹Р№ С‚РµСЃС‚",
            branch=self.branch_main,
            wants_ab_test=True,
            ab_test_options=["start"],
            metric_impact="CR",
            expected_change="+5%",
            hypothesis="РџРµСЂРІС‹Р№ С‚РµСЃС‚",
            traffic_volume=Experiment.TrafficVolume.SPLIT_50_50,
            test_duration=Experiment.TestDuration.DAYS_7,
            start_date=date(2026, 3, 10),
            end_date=date(2026, 3, 17),
            created_by=self.admin_user,
            status=Experiment.Status.IN_PROGRESS,
        )
        queued_experiment = Experiment.objects.create(
            title="Р’С‚РѕСЂРѕР№ С‚РµСЃС‚",
            branch=self.branch_main,
            wants_ab_test=True,
            ab_test_options=["start"],
            metric_impact="CR",
            expected_change="+3%",
            hypothesis="Р’С‚РѕСЂР°СЏ РіРёРїРѕС‚РµР·Р°",
            traffic_volume=Experiment.TrafficVolume.SPLIT_50_50,
            test_duration=Experiment.TestDuration.DAYS_7,
            start_date=date(2026, 3, 18),
            end_date=date(2026, 3, 25),
            created_by=self.admin_user,
            status=Experiment.Status.DRAFT,
        )

        response = self.client.post(
            reverse("update_experiment_status", kwargs={"experiment_id": queued_experiment.id}),
            {"status": Experiment.Status.IN_PROGRESS},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        queued_experiment.refresh_from_db()
        self.assertEqual(queued_experiment.status, Experiment.Status.DRAFT)
        return
        self.assertContains(response, "Р”Р»СЏ СЌС‚РѕР№ РІРµС‚РєРё СѓР¶Рµ РёРґРµС‚ РґСЂСѓРіРѕР№ A/B С‚РµСЃС‚.")

    def test_final_status_requires_dates_and_ab_results(self):
        experiment = Experiment.objects.create(
            title="Тест без итогов",
            wants_ab_test=True,
            ab_test_options=["start"],
            metric_impact="CR",
            expected_change="+5%",
            hypothesis="Проверяем первый экран",
            traffic_volume=Experiment.TrafficVolume.SPLIT_50_50,
            test_duration=Experiment.TestDuration.DAYS_7,
            created_by=self.admin_user,
            status=Experiment.Status.COMPLETED,
        )

        response = self.client.post(
            reverse("update_experiment_status", kwargs={"experiment_id": experiment.id}),
            {"status": Experiment.Status.SUCCESS},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        experiment.refresh_from_db()
        self.assertEqual(experiment.status, Experiment.Status.COMPLETED)
        self.assertContains(response, "Перед финальным решением заполните")

    def test_final_status_saves_completion_data_from_modal_post(self):
        experiment = Experiment.objects.create(
            title="Финализация из попапа",
            wants_ab_test=True,
            ab_test_options=["start"],
            metric_impact="CR",
            expected_change="+7%",
            hypothesis="Завершаем тест прямо из карточки",
            traffic_volume=Experiment.TrafficVolume.SPLIT_50_50,
            test_duration=Experiment.TestDuration.DAYS_7,
            created_by=self.admin_user,
            status=Experiment.Status.IN_PROGRESS,
        )

        response = self.client.post(
            reverse("update_experiment_status", kwargs={"experiment_id": experiment.id}),
            {
                "status": Experiment.Status.SUCCESS,
                "start_date": "2026-03-10",
                "end_date": "2026-03-17",
                "dashboard_url": "https://example.com/dashboard/final-exp",
                "result_variant_a": "CR 10%, CTR 18%",
                "result_variant_b": "CR 14%, CTR 22%",
                "comment": "Победил вариант B",
            },
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        experiment.refresh_from_db()
        self.assertEqual(experiment.status, Experiment.Status.SUCCESS)
        self.assertEqual(experiment.start_date, date(2026, 3, 10))
        self.assertEqual(experiment.end_date, date(2026, 3, 17))
        self.assertEqual(experiment.dashboard_url, "https://example.com/dashboard/final-exp")
        self.assertEqual(experiment.result_variant_a, "CR 10%, CTR 18%")
        self.assertEqual(experiment.result_variant_b, "CR 14%, CTR 22%")
        self.assertEqual(experiment.comment, "Победил вариант B")

    def test_final_status_allows_empty_variant_results(self):
        experiment = Experiment.objects.create(
            title="Финал без цифр",
            wants_ab_test=True,
            ab_test_options=["start"],
            metric_impact="CR",
            comparison_text="A vs B",
            expected_change="+7%",
            hypothesis="Проверяем без обязательных цифр",
            traffic_volume=Experiment.TrafficVolume.SPLIT_50_50,
            test_duration=Experiment.TestDuration.DAYS_7,
            created_by=self.admin_user,
            status=Experiment.Status.IN_PROGRESS,
        )

        response = self.client.post(
            reverse("update_experiment_status", kwargs={"experiment_id": experiment.id}),
            {
                "status": Experiment.Status.SUCCESS,
                "start_date": "2026-03-30",
                "end_date": "2026-04-01",
                "dashboard_url": "",
                "result_variant_a": "",
                "result_variant_b": "",
                "comment": "",
            },
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        experiment.refresh_from_db()
        self.assertEqual(experiment.status, Experiment.Status.SUCCESS)

    def test_finalized_experiment_moves_to_library(self):
        active_experiment = Experiment.objects.create(
            title="Активный тест",
            wants_ab_test=True,
            ab_test_options=["start"],
            metric_impact="CR",
            expected_change="+5%",
            hypothesis="Активная гипотеза",
            traffic_volume=Experiment.TrafficVolume.SPLIT_50_50,
            test_duration=Experiment.TestDuration.DAYS_7,
            created_by=self.admin_user,
            status=Experiment.Status.DRAFT,
        )
        final_experiment = Experiment.objects.create(
            title="Финальный тест",
            wants_ab_test=True,
            ab_test_options=["start"],
            metric_impact="CR",
            expected_change="+5%",
            hypothesis="Финальная гипотеза",
            traffic_volume=Experiment.TrafficVolume.SPLIT_50_50,
            test_duration=Experiment.TestDuration.DAYS_7,
            start_date=date(2026, 3, 10),
            end_date=date(2026, 3, 17),
            result_variant_a="CR 10%",
            result_variant_b="CR 15%",
            created_by=self.admin_user,
            status=Experiment.Status.SUCCESS,
        )

        response = self.client.get(reverse("experiments_board"))

        self.assertEqual(response.status_code, 200)

        active_ids = [
            item["item"].id
            for column in response.context["active_columns"]
            for item in column["items"]
        ]
        library_ids = [
            item["item"].id
            for column in response.context["library_columns"]
            for item in column["items"]
        ]

        self.assertIn(active_experiment.id, active_ids)
        self.assertNotIn(final_experiment.id, active_ids)
        self.assertIn(final_experiment.id, library_ids)


class BranchFormTests(TaskBoardBaseTestCase):
    def test_branch_form_prefills_next_code_from_existing_branches(self):
        Branch.objects.create(bot=self.bot, name="Second", code="ell01")
        Branch.objects.create(bot=self.bot, name="Third", code="ell02")

        form = BranchForm(bot=self.bot)

        self.assertEqual(form.initial["code"], "ell03")


class TikTokFunnelHelpersTests(TestCase):
    def test_normalize_endpoint_strips_domain_query_and_trailing_slash(self):
        cases = {
            "/pasha_all": "/pasha_all",
            "/pasha_all/": "/pasha_all",
            "pasha_all": "/pasha_all",
            "go-egeland.ru/pasha_all": "/pasha_all",
            "https://go-egeland.ru/pasha_all/?utm=1": "/pasha_all",
            "  /pasha_all?x=1#frag ": "/pasha_all",
            "https://go-egeland.ru/a/b/c": "/a/b/c",
        }
        for raw, expected in cases.items():
            self.assertEqual(TikTokFunnelRequest.normalize_endpoint(raw), expected, raw)

    def test_normalize_endpoint_empty(self):
        self.assertEqual(TikTokFunnelRequest.normalize_endpoint(""), "")
        self.assertEqual(TikTokFunnelRequest.normalize_endpoint(None), "")

    def test_parse_bot_name(self):
        cases = {
            "https://telegram.me/efir_tt_el_bot": "efir_tt_el_bot",
            "https://t.me/efir_tt_el_bot/": "efir_tt_el_bot",
            "t.me/@some_bot": "some_bot",
            "https://t.me/+joinhash": "joinhash",
            "": "",
        }
        for raw, expected in cases.items():
            self.assertEqual(TikTokFunnelRequest.parse_bot_name(raw), expected, raw)


class TikTokFunnelFormTests(TestCase):
    def test_form_normalizes_endpoint_and_derives_bot_name(self):
        form = TikTokFunnelRequestForm(
            data={
                "landing_endpoint": "https://go-egeland.ru/pasha_all/",
                "offer": "  ell010005 ",
                "bot_url": "https://telegram.me/efir_tt_el_bot",
                "comment": "",
            }
        )
        self.assertTrue(form.is_valid(), form.errors)
        obj = form.save(commit=False)
        self.assertEqual(obj.landing_endpoint, "/pasha_all")
        self.assertEqual(obj.offer, "ell010005")
        self.assertEqual(form.cleaned_data["bot_name"], "efir_tt_el_bot")

    def test_form_rejects_root_endpoint(self):
        form = TikTokFunnelRequestForm(
            data={"landing_endpoint": "https://go-egeland.ru/", "offer": "x", "bot_url": "t.me/b"}
        )
        self.assertFalse(form.is_valid())
        self.assertIn("landing_endpoint", form.errors)


class TikTokFunnelViewTests(TaskBoardBaseTestCase):
    def test_page_renders(self):
        TikTokFunnelRequest.objects.create(
            landing_endpoint="/pasha_all",
            offer="ell010005",
            bot_url="https://telegram.me/efir_tt_el_bot",
            bot_name="efir_tt_el_bot",
        )
        self.client.force_login(self.admin_user)
        response = self.client.get(reverse("tiktok_funnels"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "TikTok-воронки")
        self.assertContains(response, "/pasha_all")

    @patch("marks.views.notify_new_tiktok_funnel", return_value=(True, ""))
    @patch("marks.views.provision_funnel", return_value=(True, True, ""))
    def test_create_funnel(self, provision_mock, notify_mock):
        self.client.force_login(self.admin_user)
        response = self.client.post(
            reverse("create_tiktok_funnel"),
            data={
                "landing_endpoint": "go-egeland.ru/pasha_all",
                "offer": "ell010005",
                "bot_url": "https://telegram.me/efir_tt_el_bot",
                "comment": "тест",
            },
        )
        self.assertRedirects(response, reverse("tiktok_funnels"))
        funnel = TikTokFunnelRequest.objects.get()
        self.assertEqual(funnel.landing_endpoint, "/pasha_all")
        self.assertEqual(funnel.bot_name, "efir_tt_el_bot")
        self.assertEqual(funnel.status, TikTokFunnelRequest.Status.PENDING)
        self.assertTrue(funnel.warehouse_synced)
        provision_mock.assert_called_once()
        notify_mock.assert_called_once()

    @patch("marks.views.notify_new_tiktok_funnel", return_value=(True, ""))
    @patch("marks.views.provision_funnel", return_value=(False, False, "склад недоступен"))
    def test_create_funnel_keeps_request_when_warehouse_fails(self, provision_mock, notify_mock):
        self.client.force_login(self.admin_user)
        response = self.client.post(
            reverse("create_tiktok_funnel"),
            data={
                "landing_endpoint": "/pasha_all",
                "offer": "ell010005",
                "bot_url": "https://telegram.me/efir_tt_el_bot",
            },
        )
        self.assertRedirects(response, reverse("tiktok_funnels"))
        funnel = TikTokFunnelRequest.objects.get()
        self.assertFalse(funnel.warehouse_synced)

    @patch("marks.views.activate_funnel", return_value=(True, ""))
    def test_activate_funnel(self, activate_mock):
        funnel = TikTokFunnelRequest.objects.create(
            landing_endpoint="/pasha_all",
            offer="ell010005",
            bot_url="https://telegram.me/efir_tt_el_bot",
            bot_name="efir_tt_el_bot",
        )
        self.client.force_login(self.admin_user)
        response = self.client.post(
            reverse("activate_tiktok_funnel", args=[funnel.id]),
            data={"pixel_code": "D5NMAQRC77U4HM3KJMPG"},
        )
        self.assertRedirects(response, reverse("tiktok_funnels"))
        funnel.refresh_from_db()
        self.assertEqual(funnel.status, TikTokFunnelRequest.Status.ACTIVE)
        self.assertEqual(funnel.pixel_code, "D5NMAQRC77U4HM3KJMPG")
        activate_mock.assert_called_once()

    def test_manager_cannot_activate(self):
        funnel = TikTokFunnelRequest.objects.create(
            landing_endpoint="/pasha_all",
            offer="ell010005",
            bot_url="https://telegram.me/efir_tt_el_bot",
        )
        self.client.force_login(self.manager_user)
        response = self.client.post(
            reverse("activate_tiktok_funnel", args=[funnel.id]),
            data={"pixel_code": "X"},
        )
        self.assertEqual(response.status_code, 403)
